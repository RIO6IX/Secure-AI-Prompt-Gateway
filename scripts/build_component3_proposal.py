"""Build the standalone Component 3 proposal from its reviewable Markdown source."""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "COMPONENT3_PROPOSAL.md"
OUTPUT = ROOT / "docs" / "Shaheed_Component3_Proposal.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def borders(table):
    props = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        node.append(border)
    props.append(node)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def keep_row_together(row):
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))


def field(paragraph, instruction):
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t"); fallback.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    for item in (begin, text, separate, fallback, end): run.append(item)


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part: continue
        if part.startswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9.5)
        elif part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def parse_table(lines, start):
    rows, index = [], start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value) for value in cells): rows.append(cells)
        index += 1
    return rows, index


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.4)
    section.left_margin = section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12); style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Security Policy and Access Control Engine")
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Component 3 Proposal and Technical Specification").bold = True
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Secure AI Prompt Gateway\nIE3092 Information Security Project\n\nH. M. Shaheed   IT23664258\nSri Lanka Institute of Information Technology\n9 September 2026")
    intro = doc.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run("This document specifies the implemented authentication, role-based access control, versioned policy evaluation and approval workflow owned by Component 3. It also states the verified behavior and the boundaries that remain for group integration.")
    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    for line in lines:
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="List Bullet")
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="List Bullet 2")
            paragraph.add_run(line[4:].strip())
    doc.add_page_break()

    index = 0
    while index < len(lines):
        raw = lines[index].strip()
        if not raw or raw.startswith("# ") or raw.startswith("Owner:") or raw.startswith("Project:") or raw.startswith("Implementation branch:") or raw.startswith("Revision:"):
            index += 1; continue
        if raw.startswith("### "):
            doc.add_heading(raw[4:], level=2); index += 1; continue
        if raw.startswith("## "):
            doc.add_heading(raw[3:], level=1); index += 1; continue
        if raw.startswith("|"):
            rows, index = parse_table(lines, index)
            if not rows: continue
            table = doc.add_table(rows=len(rows), cols=len(rows[0])); table.autofit = True; borders(table)
            repeat_header(table.rows[0])
            for row in table.rows:
                keep_row_together(row)
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    cell = table.cell(r, c); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = value
                    for run in cell.paragraphs[0].runs:
                        run.font.name = "Arial"; run.font.size = Pt(8.5); run.bold = r == 0
                        if r == 0: run.font.color.rgb = RGBColor(255, 255, 255)
                    if r == 0: shade(cell, "17365D")
                    elif r % 2 == 0: shade(cell, "EAF2F8")
            doc.add_paragraph()
            continue
        if re.match(r"^\d+\. ", raw):
            paragraph = doc.add_paragraph(style="List Number"); add_inline(paragraph, re.sub(r"^\d+\. ", "", raw)); index += 1; continue
        paragraph = doc.add_paragraph(); add_inline(paragraph, raw); index += 1

    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("H. M. Shaheed   Component 3   |   Page "); field(footer, "PAGE")
    settings = doc.settings.element
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    doc.core_properties.title = "Security Policy and Access Control Engine"
    doc.core_properties.subject = "IE3092 Secure AI Prompt Gateway Component 3"
    doc.core_properties.author = "H. M. Shaheed"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__": main()
